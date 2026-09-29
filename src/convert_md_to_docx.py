import os
import re
import sys
import hashlib
import urllib.request
import base64
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml

# Setup MML2OMML XSLT for native Word Equation conversion
XSL_PATHS = [
    r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL',
    r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL'
]
OMML_TRANSFORM = None
for xp in XSL_PATHS:
    if os.path.exists(xp):
        try:
            xslt_doc = etree.parse(xp)
            OMML_TRANSFORM = etree.XSLT(xslt_doc)
            break
        except Exception:
            pass

def latex_to_omml_element(latex_str):
    if OMML_TRANSFORM is None:
        return None
    try:
        clean = latex_str.strip()
        if not clean:
            return None
        # Handle some common latex replacements if needed
        clean = clean.replace(r'\ge', r'\ge ').replace(r'\le', r'\le ')
        mathml = latex_to_mathml(clean)
        mathml_tree = etree.fromstring(mathml)
        omml_tree = OMML_TRANSFORM(mathml_tree)
        omml_bytes = etree.tostring(omml_tree, encoding='utf-8')
        elem = parse_xml(omml_bytes.decode('utf-8'))
        return elem
    except Exception:
        return None

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

INLINE_PATTERN = re.compile(
    r'(?P<bold_italic>\*\*\*(?P<bi_text>.*?)\*\*\*)|'
    r'(?P<bold>\*\*(?P<b_text>.*?)\*\*)|'
    r'(?P<italic>\*(?P<it_text>.*?)\*)|'
    r'(?P<code>`(?P<code_text>.*?)`)|'
    r'(?P<math>\$(?P<math_text>[^$]+?)\$)|'
    r'(?P<link>\[(?P<link_text>.*?)\]\((?P<link_url>.*?)\))'
)

def parse_inline_markdown(paragraph, text, base_font_size=Pt(11), base_color=RGBColor(0x2D, 0x37, 0x48)):
    last_idx = 0
    for match in INLINE_PATTERN.finditer(text):
        start, end = match.span()
        if start > last_idx:
            run = paragraph.add_run(text[last_idx:start])
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            
        m_dict = match.groupdict()
        if m_dict.get('bold_italic'):
            run = paragraph.add_run(m_dict['bi_text'])
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.bold = True
            run.italic = True
        elif m_dict.get('bold'):
            run = paragraph.add_run(m_dict['b_text'])
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.bold = True
        elif m_dict.get('italic'):
            run = paragraph.add_run(m_dict['it_text'])
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.italic = True
        elif m_dict.get('code'):
            run = paragraph.add_run(m_dict['code_text'])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x9C, 0x1A, 0x1A)
        elif m_dict.get('math'):
            math_text = m_dict['math_text']
            omml_elem = latex_to_omml_element(math_text)
            if omml_elem is not None:
                paragraph._p.append(omml_elem)
            else:
                # Fallback to styled text
                run = paragraph.add_run(math_text)
                run.font.name = 'Cambria Math'
                run.font.size = base_font_size
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
                run.italic = True
        elif m_dict.get('link'):
            run = paragraph.add_run(m_dict['link_text'])
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            run.underline = True
            
        last_idx = end
        
    if last_idx < len(text):
        run = paragraph.add_run(text[last_idx:])
        run.font.name = 'Calibri'
        run.font.size = base_font_size
        run.font.color.rgb = base_color

def render_mermaid_to_png(mermaid_code, cache_dir='docs/assets/mermaid_cache'):
    os.makedirs(cache_dir, exist_ok=True)
    clean_code = mermaid_code.strip()
    code_hash = hashlib.md5(clean_code.encode('utf-8')).hexdigest()
    cache_path = os.path.join(cache_dir, f"{code_hash}.png")
    
    if os.path.exists(cache_path) and os.path.getsize(cache_path) > 500:
        return cache_path
        
    try:
        graphbytes = clean_code.encode('utf-8')
        base64_bytes = base64.b64encode(graphbytes)
        base64_string = base64_bytes.decode('ascii')
        url = 'https://mermaid.ink/img/' + base64_string
        
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
        with urllib.request.urlopen(req, timeout=12) as response:
            img_data = response.read()
            if len(img_data) > 500:
                with open(cache_path, 'wb') as f:
                    f.write(img_data)
                return cache_path
    except Exception:
        pass
        
    return None

def convert_md_file(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = docx.Document()
    
    # Page setup: Standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    md_dir = os.path.dirname(md_path)
    workspace_root = r"e:\Project Buku"
    
    in_code_block = False
    code_lang = ""
    code_block_lines = []
    in_table = False
    table_lines = []
    in_math_block = False
    math_block_lines = []
    
    def flush_table():
        nonlocal table_lines, in_table
        if not table_lines:
            in_table = False
            return
        
        rows_data = []
        for tl in table_lines:
            tl_strip = tl.strip()
            if not tl_strip.startswith('|'):
                continue
            cells = [c.strip() for c in tl_strip.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            rows_data.append(cells)
            
        if rows_data:
            n_rows = len(rows_data)
            n_cols = max(len(r) for r in rows_data)
            
            table = doc.add_table(rows=n_rows, cols=n_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)
            
            for row_idx, rdata in enumerate(rows_data):
                row = table.rows[row_idx]
                is_header = (row_idx == 0)
                for col_idx in range(n_cols):
                    cell = row.cells[col_idx]
                    cell_text = rdata[col_idx] if col_idx < len(rdata) else ""
                    
                    if is_header:
                        set_cell_background(cell, "1A365D")
                    elif row_idx % 2 == 1:
                        set_cell_background(cell, "F7FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                        
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.15
                    
                    if is_header:
                        run = cp.add_run(cell_text)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        parse_inline_markdown(cp, cell_text, base_font_size=Pt(9.5))
                        
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            
        table_lines = []
        in_table = False

    def flush_code_block():
        nonlocal code_block_lines, in_code_block, code_lang
        if not code_block_lines:
            in_code_block = False
            code_lang = ""
            return
            
        code_text = "".join(code_block_lines)
        
        # Check if code block is Mermaid diagram!
        if code_lang == 'mermaid' or code_text.strip().startswith(('flowchart', 'graph', 'sequenceDiagram', 'classDiagram', 'stateDiagram')):
            img_path = render_mermaid_to_png(code_text)
            if img_path and os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.5))
                
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run("Diagram Alir Konseptual (Visualisasi Mermaid)")
                run_cap.font.name = 'Calibri'
                run_cap.font.size = Pt(9)
                run_cap.font.italic = True
                run_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                
                code_block_lines = []
                in_code_block = False
                code_lang = ""
                return
        
        # Standard Code Block
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
            f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="3182CE"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(code_text.rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        code_block_lines = []
        in_code_block = False
        code_lang = ""

    def render_display_equation(math_code):
        p_eq = doc.add_paragraph()
        p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_eq.paragraph_format.space_before = Pt(6)
        p_eq.paragraph_format.space_after = Pt(6)
        omml_elem = latex_to_omml_element(math_code)
        if omml_elem is not None:
            p_eq._p.append(omml_elem)
        else:
            run = p_eq.add_run(math_code)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(12)
            run.italic = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    i = 0
    while i < len(lines):
        line = lines[i]
        line_strip = line.strip()
        
        # Check display math block $$ ... $$
        if line_strip == '$$':
            if in_math_block:
                render_display_equation("\n".join(math_block_lines))
                math_block_lines = []
                in_math_block = False
            else:
                in_math_block = True
                math_block_lines = []
            i += 1
            continue
            
        if in_math_block:
            math_block_lines.append(line_strip)
            i += 1
            continue
            
        # Single line display math: $$ formula $$
        single_math_match = re.match(r'^\$\$(.*?)\$\$$', line_strip)
        if single_math_match:
            render_display_equation(single_math_match.group(1))
            i += 1
            continue
        
        # Check code fence
        if line_strip.startswith('```'):
            if in_table:
                flush_table()
            if in_code_block:
                flush_code_block()
            else:
                in_code_block = True
                code_lang = line_strip[3:].strip().lower()
                code_block_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Check table
        if line_strip.startswith('|'):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()
            
        # Check image: ![caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line_strip)
        if img_match:
            caption, img_rel = img_match.group(1), img_match.group(2)
            img_full = os.path.normpath(os.path.join(md_dir, img_rel))
            if not os.path.exists(img_full):
                asset_name = os.path.basename(img_rel)
                img_full = os.path.join(workspace_root, 'docs', 'assets', asset_name)
                
            if os.path.exists(img_full):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                run_img.add_picture(img_full, width=Inches(5.8))
                
                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(8)
                    run_cap = p_cap.add_run(f"Gambar: {caption}")
                    run_cap.font.name = 'Calibri'
                    run_cap.font.size = Pt(9)
                    run_cap.font.italic = True
                    run_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
            i += 1
            continue
            
        # Heading 1
        if line_strip.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_strip[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            i += 1
            continue
            
        # Heading 2
        if line_strip.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_strip[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            i += 1
            continue
            
        # Heading 3
        if line_strip.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_strip[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            i += 1
            continue
            
        # Heading 4
        if line_strip.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line_strip[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            i += 1
            continue
            
        # Horizontal Rule
        if line_strip in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE0)
            run.font.size = Pt(8)
            i += 1
            continue
            
        # Blockquote / Callout (> ...)
        if line_strip.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            bq_text = re.sub(r'^>\s*(\[!.*?\])?\s*', '', line_strip)
            run_bar = p.add_run("┃ ")
            run_bar.font.color.rgb = RGBColor(0x31, 0x82, 0xCE)
            run_bar.font.bold = True
            parse_inline_markdown(p, bq_text, base_font_size=Pt(10), base_color=RGBColor(0x4A, 0x55, 0x68))
            i += 1
            continue
            
        # Bullet list (- or *)
        if re.match(r'^\s*[-*]\s+', line):
            indent_level = len(re.match(r'^\s*', line).group(0)) // 2
            bullet_text = re.sub(r'^\s*[-*]\s+', '', line).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_markdown(p, bullet_text)
            i += 1
            continue
            
        # Numbered list (1. or 2.)
        num_match = re.match(r'^\s*(\d+)\.\s+(.*)$', line_strip)
        if num_match:
            num, item_text = num_match.group(1), num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run_num = p.add_run(f"{num}. ")
            run_num.font.name = 'Calibri'
            run_num.font.size = Pt(11)
            run_num.font.bold = True
            run_num.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            parse_inline_markdown(p, item_text)
            i += 1
            continue
            
        # Normal Paragraph
        if line_strip:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_markdown(p, line_strip)
            
        i += 1
        
    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()
        
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    return docx_path

def main(target_parts=None):
    if target_parts:
        parts = target_parts
    else:
        parts = [p for p in sorted(os.listdir('docs')) if p.startswith('part-')]
    total_converted = 0
    errors = []
    
    print("==========================================================")
    print("MEMULAI PROSES KONVERSI BATCH MARKDOWN KE DOCX")
    print("DENGAN PERSAMAAN MATEMATIKA NATIVE WORD (OMML) & MERMAID")
    print(f"Parts: {', '.join(parts)}")
    print("==========================================================")
    
    for part in parts:
        in_folder = os.path.join('docs', part)
        if not os.path.exists(in_folder):
            continue
        out_folder = os.path.join('docx', part)
        os.makedirs(out_folder, exist_ok=True)
        
        md_files = [f for f in sorted(os.listdir(in_folder)) if f.endswith('.md')]
        print(f"\n[+] Memproses {part.upper()} ({len(md_files)} modul)...")
        
        for f in md_files:
            src = os.path.join(in_folder, f)
            dst_name = f.replace('.md', '.docx')
            dst = os.path.join(out_folder, dst_name)
            
            try:
                convert_md_file(src, dst)
                size_kb = os.path.getsize(dst) / 1024
                print(f"  [OK] {dst_name} ({size_kb:.1f} KB)")
                total_converted += 1
            except Exception as e:
                print(f"  [FAIL] {f} -> {e}")
                errors.append((f, str(e)))
                
    # Memproses Instructor Resources
    for part in parts:
        in_folder = os.path.join('instructor_resources', part)
        if not os.path.exists(in_folder):
            continue
        out_folder_inst = in_folder
        out_folder_docx = os.path.join('docx', 'instructor_resources', part)
        os.makedirs(out_folder_docx, exist_ok=True)
        
        md_files = [f for f in sorted(os.listdir(in_folder)) if f.endswith('.md')]
        print(f"\n[+] Memproses INSTRUCTOR RESOURCES {part.upper()} ({len(md_files)} modul)...")
        
        for f in md_files:
            src = os.path.join(in_folder, f)
            dst_name = f.replace('.md', '.docx')
            dst1 = os.path.join(out_folder_inst, dst_name)
            dst2 = os.path.join(out_folder_docx, dst_name)
            
            try:
                convert_md_file(src, dst1)
                import shutil
                shutil.copyfile(dst1, dst2)
                size_kb = os.path.getsize(dst1) / 1024
                print(f"  [OK] {dst_name} ({size_kb:.1f} KB)")
                total_converted += 1
            except Exception as e:
                print(f"  [FAIL] {f} -> {e}")
                errors.append((f, str(e)))

    print("\n==========================================================")
    print(f"SELESAI: {total_converted} file DOCX berhasil diproses!")
    if errors:
        print(f"Terdapat {len(errors)} error:")
        for err in errors:
            print(f" - {err[0]}: {err[1]}")
    else:
        print("Status: 100% SUKSES TANPA ERROR!")
    print("==========================================================")

if __name__ == '__main__':
    if len(sys.argv) == 3 and not sys.argv[1].startswith('part-'):
        convert_md_file(sys.argv[1], sys.argv[2])
        print(f"[OK] Berhasil mengonversi: {sys.argv[1]} -> {sys.argv[2]}")
    elif len(sys.argv) > 1 and sys.argv[1].startswith('part-'):
        main(sys.argv[1:])
    else:
        main()

