import glob
import os

pptx_files = sorted(glob.glob('pptx/part-*/*.pptx'))
slide_files = sorted(glob.glob('slides/part-*/*.md'))

print(f'Total PPTX files: {len(pptx_files)}')
print(f'Total Slide MD files: {len(slide_files)}')

banned = ['taksonomi', 'tripartit', 'ontologis', 'epistemologis', 'dialektika', 'manifestasi dependensi', 'de facto']
found_banned = 0
for f in slide_files:
    content = open(f, encoding='utf-8').read().lower()
    for b in banned:
        if b in content:
            print(f'Found banned word "{b}" in {f}')
            found_banned += 1

print(f'Total banned words in slides: {found_banned}')

# Also verify PPTX files can be loaded
from pptx import Presentation
load_errors = 0
for f in pptx_files:
    try:
        prs = Presentation(f)
        if len(prs.slides) != 8:
            print(f'Warning: {f} has {len(prs.slides)} slides instead of 8')
    except Exception as e:
        print(f'Error reading {f}: {e}')
        load_errors += 1

print(f'Total PPTX load errors: {load_errors}')
print('Verification completed successfully!')
